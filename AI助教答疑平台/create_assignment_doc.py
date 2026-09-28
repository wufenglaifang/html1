from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


def set_font(run, size=None, bold=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_bullet(document, text):
    p = document.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    set_font(run, 10.5)


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.3)
section.bottom_margin = Cm(2.3)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.5)

normal = doc.styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("AI助教答疑平台作业说明")
set_font(run, 20, True)
title.paragraph_format.space_after = Pt(14)

intro = doc.add_paragraph()
intro.paragraph_format.space_after = Pt(13)
run = intro.add_run("项目内容：")
set_font(run, 10.5, True)
run = intro.add_run("使用原生 HTML5 和 CSS 完成 AI 助教答疑平台的登录、注册页面。页面采用淡蓝渐变背景、居中白色圆角卡片和蓝色按钮，并适配移动端显示。")
set_font(run, 10.5)

heading1 = doc.add_paragraph(style="Heading 1")
heading1.paragraph_format.space_before = Pt(4)
heading1.paragraph_format.space_after = Pt(8)
run = heading1.add_run("一、用到的 HTML5 表单增强特性")
set_font(run, 14, True)

features = [
    "使用 text、password、email、tel、date、checkbox 等 input 类型，分别接收用户名、密码、邮箱、手机号、出生日期和协议确认信息。",
    "使用 required 属性实现必填校验，未填写完整时浏览器会自动提示。",
    "使用 pattern 正则表达式校验用户名、邮箱和中国大陆手机号码格式。",
    "使用 minlength、maxlength 限制密码长度为 6 至 20 位。",
    "使用 placeholder 为输入框提供填写提示。",
    "使用 autocomplete 支持浏览器自动填充账号、密码、邮箱、手机号和出生日期。",
    "登录页面使用 autofocus，打开页面时自动聚焦用户名输入框。",
    "使用 label 的 for 属性与输入框 id 绑定，提升表单可用性。",
    "使用 fieldset 和 legend 对登录、注册表单进行分组，使结构更清晰。",
]
for item in features:
    add_bullet(doc, item)

heading2 = doc.add_paragraph(style="Heading 1")
heading2.paragraph_format.space_before = Pt(10)
heading2.paragraph_format.space_after = Pt(8)
run = heading2.add_run("二、用到的其他 HTML5 相关能力")
set_font(run, 14, True)

other = [
    "使用 header 放置平台图标、页面标题和副标题。",
    "使用 section 放置登录或注册表单卡片。",
    "使用 footer 放置“前往注册”和“去登录”的页面跳转链接。",
    "使用 meta viewport 实现移动端自适应。",
    "使用 CSS3 的渐变背景、圆角、阴影、弹性布局和媒体查询，实现统一、简约的页面视觉效果。",
    "登录页与注册页通过超链接相互跳转，形成基本的用户入口流程。",
]
for item in other:
    add_bullet(doc, item)

doc.save("AI助教答疑平台作业说明.docx")
